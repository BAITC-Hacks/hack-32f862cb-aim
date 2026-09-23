"""Read-only adapters for the supplied IEK/Systeme workbooks; never execute Excel formulas."""

import hashlib
import math
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from optistock.config import settings
from optistock.errors import DomainError

MONTHS = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]


def text(value):
    return "" if value is None else str(value).strip()


def number(value):
    if value is None or text(value) == "":
        return None
    try:
        result = float(text(value).replace("\xa0", "").replace(" ", "").replace(",", "."))
        if not math.isfinite(result) or abs(result) > 1e12:
            raise ValueError
        return result
    except ValueError:
        raise DomainError("invalid_number", "Некорректное числовое значение в Excel") from None


def month(value):
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-01")
    v = text(value).lower()
    year = re.search(r"20\d{2}", v)
    if not year:
        return None
    for i, name in enumerate(MONTHS, 1):
        if name in v:
            return f"{year.group()}-{i:02d}-01"
    return None


def excel_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text(value), fmt).date()
        except ValueError:
            continue
    raise DomainError("invalid_date", "Некорректная дата в Excel")


def validate_xlsx(path: Path):
    try:
        with ZipFile(path) as archive:
            files = archive.infolist()
            if len(files) > 2000 or sum(f.file_size for f in files) > settings().xlsx_max_uncompressed_bytes:
                raise DomainError("xlsx_too_large", "Превышен лимит распакованного Excel", 413)
            if "xl/workbook.xml" not in archive.namelist():
                raise DomainError("invalid_xlsx", "Ожидается файл XLSX")
            if any(f.filename.lower().endswith("vbaproject.bin") for f in files):
                raise DomainError("macros_not_allowed", "Макросы не поддерживаются")
    except BadZipFile:
        raise DomainError("invalid_xlsx", "Файл не является корректным XLSX") from None


def store_file(stream) -> str:
    root = settings().storage_root
    root.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256()
    size = 0
    target = tempfile.NamedTemporaryFile(dir=root, delete=False)
    temporary = Path(target.name)
    try:
        with target:
            while chunk := stream.read(1024 * 1024):
                size += len(chunk)
                if size > settings().upload_max_bytes:
                    raise DomainError("file_too_large", "Файл превышает лимит загрузки", 413)
                digest.update(chunk)
                target.write(chunk)
            target.flush()
            os.fsync(target.fileno())
        # Windows cannot rename or unlink an open file. Close it before validation/publication.
        validate_xlsx(temporary)
        key = digest.hexdigest()
        if not (root / f"{key}.xlsx").exists():
            temporary.replace(root / f"{key}.xlsx")
        return key
    finally:
        temporary.unlink(missing_ok=True)


@dataclass
class ParsedItem:
    supplier: str
    code: str
    name: str = ""
    unit: str = "шт"
    category: str | None = None
    minimum: float | None = None
    multiple: float | None = None
    available: float | None = None
    data: dict = field(
        default_factory=lambda: {
            "sales": {},
            "stocks": {},
            "events": [],
            "incoming": [],
            "warnings": [],
            "sources": {},
        }
    )


def parse_sources(sources: list[dict], as_of: date, progress=lambda **kwargs: None):
    items: dict[tuple[str, str], ParsedItem] = {}
    seasonal: dict[str, list] = {}
    kinds = []

    def item(supplier, code, name=""):
        code = text(code)
        if not code or code.lower() in {"итого", "код", "номенклатура.код"}:
            return None
        key = (supplier, code)
        if key not in items:
            items[key] = ParsedItem(supplier=supplier, code=code, name=text(name))
        elif name and not items[key].name:
            items[key].name = text(name)
        return items[key]

    for file_index, source in enumerate(sources):
        path = settings().storage_root / f"{source['sha256']}.xlsx"
        if hashlib.sha256(path.read_bytes()).hexdigest() != source["sha256"]:
            raise DomainError("source_changed", "Контрольная сумма исходного файла не совпадает")
        validate_xlsx(path)
        supplier = source["supplier"]
        progress(stage="reading", file=file_index + 1, total_files=len(sources), name=source["name"])
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book["TDSheet"] if "TDSheet" in book.sheetnames else book.worksheets[0]
            if (
                (sheet.max_column or 0) > 256
                or (sheet.max_row or 0) > settings().xlsx_max_rows
                or (sheet.max_row or 0) * (sheet.max_column or 0) > 20_000_000
            ):
                raise DomainError("sheet_too_large", "Превышен лимит размеров таблицы", 413)
            rows = sheet.iter_rows(values_only=True)
            head = [next(rows, ()) for _ in range(3)]
            headers = [text(v).lower() for r in head for v in r]
            filename = source["name"].lower()
            if "дата" in headers and "документ" in headers:
                kind, header_row = "transactions", 0
            elif "свободный остаток" in headers:
                kind, header_row = "systeme_transit", 1
            elif any("поступление до" in h for h in headers):
                kind, header_row = "iek_transit", 0
            elif any("мин. разр" in h or "кратность" == h or "миним" in h for h in headers) and not any(
                month(v) for v in head[0]
            ):
                kind, header_row = "moq", 0
            elif any(month(v) for v in head[0]):
                if "остат" in filename:
                    kind, header_row = "stocks", 0
                elif "продаж" in filename:
                    kind, header_row = "sales", 0
                else:
                    raise DomainError("ambiguous_schema", "Не определён тип месячной таблицы")
            elif "год" in headers:
                kind, header_row = "seasonality", 2
            else:
                raise DomainError("unknown_schema", f"Не распознана схема: {source['name']}")
            kinds.append({"id": source["id"], "kind": kind})
            header = [text(v).lower() for v in head[header_row]]
            month_columns = {i: month(v) for i, v in enumerate(head[header_row]) if month(v)}
            code_col = next(
                (i for i, h in enumerate(header) if h in ("код", "код 1с", "номенклатура.код")), None
            )
            name_col = next((i for i, h in enumerate(header) if h in ("номенклатура", "наименование")), None)
            unit_col = next((i for i, h in enumerate(header) if h in ("ед.", "ед.изм")), None)
            moq_col = next((i for i, h in enumerate(header) if "мин" in h or h == "кратность"), None)
            seen = set()
            # Restart iterator so the first data rows buffered above are never lost.
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if row_no <= header_row + 1:
                    continue
                if row_no > settings().xlsx_max_rows:
                    raise DomainError("too_many_rows", "Превышен лимит строк Excel", 413)
                ref = {"file_id": source["id"], "sheet": sheet.title, "row": row_no}
                if kind == "seasonality":
                    if isinstance(row[0], (int, float)) and 2000 <= row[0] <= as_of.year:
                        vals = [number(v) for v in row[1:13]]
                        if (
                            len(vals) == 12
                            and all(v is not None and v >= 0 for v in vals)
                            and int(row[0]) < as_of.year
                        ):
                            seasonal.setdefault(supplier, []).append(
                                {"year": int(row[0]), "values": vals, "source": ref}
                            )
                    continue
                if code_col is None or name_col is None:
                    raise DomainError("missing_columns", f"Нет кода или наименования: {source['name']}")
                p = item(supplier, row[code_col], row[name_col])
                if p is None:
                    continue
                if unit_col is not None and row[unit_col]:
                    p.unit = text(row[unit_col])
                if p.code in seen and kind not in {"transactions"}:
                    p.data["warnings"].append(f"duplicate_{kind}_row")
                seen.add(p.code)
                p.data["sources"][kind] = ref
                if kind in {"sales", "stocks"}:
                    target = p.data[kind]
                    for col, period in month_columns.items():
                        if period > as_of.isoformat():
                            continue
                        v = number(row[col])
                        if period in target and target[period] != v:
                            raise DomainError(
                                "conflicting_month", f"Конфликт месячных данных: {supplier}/{p.code}"
                            )
                        target[period] = v
                elif kind == "moq":
                    try:
                        q = number(row[moq_col]) if moq_col is not None else None
                    except DomainError:
                        q = None
                        p.data["raw_invalid_moq"] = text(row[moq_col])
                    if q is not None and q > 0:
                        p.minimum = max(p.minimum or 0, q)
                        if "кратность" in header:
                            p.multiple = max(p.multiple or 0, q)
                    else:
                        p.data["warnings"].append("invalid_moq_default_used")
                elif kind == "transactions":
                    d, q = excel_date(row[0]), number(row[7])
                    if q is None:
                        p.data["warnings"].append("transaction_quantity_missing")
                        p.data.setdefault("skipped_transaction_rows", []).append(ref)
                    if d <= as_of and q is not None:
                        p.data["events"].append(
                            {"date": d.isoformat(), "q": q, "document": text(row[1]), "source": ref}
                        )
                elif kind == "iek_transit":
                    if "БУХТАМИ" in text(row[name_col]).upper():
                        p.data["cable_reels"] = True
                    for col, h in enumerate(header):
                        matches = re.findall(r"до (\d{2}\.\d{2}\.\d{4})", h)
                        if matches:
                            q = number(row[col])
                            if q is not None and q > 0:
                                p.data["incoming"].append(
                                    {
                                        "eta": excel_date(matches[-1]).isoformat(),
                                        "q": q,
                                        "source": {**ref, "column": col + 1},
                                    }
                                )
                elif kind == "systeme_transit":
                    p.category = text(row[4]) or None
                    p.available = number(row[51])
                    p.data["inventory_source"] = {**ref, "column": 52}
                    q = number(row[54])
                    if q is not None and q > 0:
                        match = re.search(r"(\d{2})\.(\d{2})", header[54])
                        if match is None:
                            raise DomainError("invalid_eta", "Не найдена дата поставки Systeme Electric")
                        eta = date(as_of.year, int(match[2]), int(match[1]))
                        p.data["incoming"].append(
                            {"eta": eta.isoformat(), "q": q, "source": {**ref, "column": 55}}
                        )
        finally:
            book.close()
    for p in items.values():
        p.data["seasonality"] = seasonal.get(p.supplier, [])
        p.data["warnings"] = sorted(set(p.data["warnings"]))
    if not items:
        raise DomainError("empty_dataset", "Не найдено ни одного товара")
    return list(items.values()), kinds
